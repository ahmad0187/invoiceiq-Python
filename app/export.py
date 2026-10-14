"""CSV and Excel exports (one row per invoice plus a line-items sheet)."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Invoice

INVOICE_COLUMNS = ["id", "vendor", "invoice_number", "invoice_date", "due_date", "bill_to", "currency",
                   "subtotal", "tax", "total", "status", "issues"]


def _row(inv: Invoice) -> list:
    return [inv.id, inv.vendor, inv.invoice_number, inv.invoice_date, inv.due_date, inv.bill_to, inv.currency,
            inv.subtotal, inv.tax, inv.total, inv.status, "; ".join(i["message"] for i in inv.issues or [])]


def to_csv(invoices: list[Invoice]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(INVOICE_COLUMNS)
    for inv in invoices:
        w.writerow(["" if v is None else v for v in _row(inv)])
    return buf.getvalue()


def to_xlsx(invoices: list[Invoice]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([c.replace("_", " ").title() for c in INVOICE_COLUMNS])
    for inv in invoices:
        ws.append(_row(inv))

    items = wb.create_sheet("Line items")
    items.append(["Invoice ID", "Vendor", "Invoice Number", "Description", "Quantity", "Unit Price", "Amount"])
    for inv in invoices:
        for li in inv.line_items:
            items.append([inv.id, inv.vendor, inv.invoice_number, li.description, li.quantity, li.unit_price, li.amount])

    header_fill = PatternFill("solid", fgColor="1F2937")
    for sheet, money_cols, date_cols in ((ws, (8, 9, 10), (4, 5)), (items, (6, 7), ())):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for c in money_cols:
                row[c - 1].number_format = "#,##0.00"
            for c in date_cols:
                row[c - 1].number_format = "yyyy-mm-dd"
        for i, col in enumerate(sheet.columns, start=1):
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        sheet.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
