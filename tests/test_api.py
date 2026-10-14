from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import sample


def upload(client, *names):
    files = [("files", (sample(n).name, sample(n).read_bytes(), "application/pdf")) for n in names]
    return client.post("/api/invoices", files=files)


def test_health(client):
    assert client.get("/api/health").json() == {"ok": True, "extractor": "rules"}


def test_upload_extracts_and_validates(client):
    res = upload(client, "01", "04")
    assert res.status_code == 201
    good, bad = res.json()
    assert good["status"] == "ready" and good["total"] == 198.72 and len(good["line_items"]) == 3
    assert bad["status"] == "needs_review"
    assert [i["code"] for i in bad["issues"]] == ["total_mismatch"]


def test_rejects_non_pdf(client):
    res = client.post("/api/invoices", files=[("files", ("notes.pdf", b"hello", "application/pdf"))])
    assert res.status_code == 400


def test_duplicate_detection(client):
    upload(client, "02")
    second = upload(client, "02").json()[0]
    assert "duplicate" in [i["code"] for i in second["issues"]]


def test_fix_then_approve(client):
    bad = upload(client, "04").json()[0]
    assert client.post(f"/api/invoices/{bad['id']}/approve").status_code == 409
    fixed = client.patch(f"/api/invoices/{bad['id']}", json={"total": 658.32}).json()
    assert fixed["issues"] == [] and fixed["status"] == "ready"
    approved = client.post(f"/api/invoices/{bad['id']}/approve").json()
    assert approved["status"] == "approved"


def test_filters_and_search(client):
    upload(client, "01", "03", "04")
    assert len(client.get("/api/invoices?status=needs_review").json()) == 1
    assert [i["vendor"] for i in client.get("/api/invoices?q=northwind").json()] == ["Northwind Office Supplies Inc."]


def test_exports(client):
    upload(client, "01", "02")
    csv = client.get("/api/export?format=csv").text.splitlines()
    assert csv[0].startswith("id,vendor,invoice_number") and len(csv) == 3
    wb = load_workbook(BytesIO(client.get("/api/export?format=xlsx").content))
    assert wb.sheetnames == ["Invoices", "Line items"]
    assert wb["Line items"].max_row == 1 + 3 + 4


def test_delete(client):
    inv = upload(client, "01").json()[0]
    assert client.delete(f"/api/invoices/{inv['id']}").status_code == 204
    assert client.get(f"/api/invoices/{inv['id']}").status_code == 404


def test_dashboard_page(client):
    res = client.get("/")
    assert res.status_code == 200 and "InvoiceIQ" in res.text
